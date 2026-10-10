"""Atomic Excel and paginated Unicode table-PDF timesheets."""

from datetime import timedelta
from html import escape
import math
from pathlib import Path

from worklogger.domain.shared.errors import InfrastructureError, ValidationError
from worklogger.domain.shared.result import Result
from worklogger.infrastructure.files import atomic_destination
from worklogger.infrastructure.i18n import _
from worklogger.infrastructure.work_type_labels import work_type_label


def _headers():
    return (_("Date"), _("Start"), _("End"), _("Work type"), _("Project"), _("Work item"), _("Content"),
            _("Worked hours"), _("Rest hours"), _("Leave hours"), _("Record status"))


def _rows(sheet):
    by_day = {}
    for entry in sheet.entries:
        by_day.setdefault(entry.day, []).append(entry)
    statuses = {"recorded": _("Recorded"), "scheduled": _("Scheduled"), "historical": _("Historical")}
    for index in range((sheet.end - sheet.start).days + 1):
        day = sheet.start + timedelta(days=index)
        entries = sorted(by_day.get(day, ()), key=lambda entry: (entry.start_time or "", entry.id or 0))
        if not entries:
            yield (day.isoformat(), "", "", "", "", "", "", 0.0, 0.0, 0.0, _("No records")), None
        for entry in entries:
            rest = entry.raw_hours() if entry.is_break else entry.break_hours
            yield (day.isoformat(), entry.start_time or "", entry.end_time or "", work_type_label(entry.work_type),
                entry.context.project_label, entry.context.work_item_label, entry.note, entry.worked_hours(), rest,
                entry.leave_hours(standard_hours=sheet.standard_hours), statuses[sheet.record_status(entry)]), entry


def _summary(sheet):
    stats = sheet.statistics
    return ((_("Worked hours"), stats.total_hours), (_("Rest hours"), sheet.rest_hours),
            (_("Leave hours"), sheet.leave_hours), (_("Overtime hours"), stats.overtime_hours))


class TimesheetXlsxExporter:
    def export_timesheet(self, destination, sheet):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
            workbook = Workbook()
            page = workbook.active
            page.title = _("Timesheet")[:31]
            page.merge_cells("A1:K1")
            page["A1"] = _("Timesheet")
            page["A1"].font = Font(name="Arial", size=18, bold=True, color="1E2035")
            page.merge_cells("A2:K2")
            page["A2"] = f"{sheet.start.isoformat()} - {sheet.end.isoformat()}"
            page.merge_cells("A3:K3")
            _text(page["A3"], _("Prepared by: {name}").format(name=sheet.author))
            page.merge_cells("A4:K4")
            page["A4"] = _("Generated at: {time}").format(time=sheet.generated_at.isoformat(timespec="seconds"))
            for column, value in enumerate(_headers(), 1):
                cell = page.cell(6, column, value)
                cell.fill = PatternFill("solid", fgColor="E8F0FE")
                cell.font = Font(name="Arial", bold=True, color="1E2035")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            border = Border(bottom=Side(style="hair", color="CDD5F0"))
            record_font = Font(name="Arial", size=10, color="1E2035")
            record_alignment = Alignment(wrap_text=True, vertical="top")
            row_index = 7
            metadata = []
            for values, entry in _rows(sheet):
                for column, value in enumerate(values, 1):
                    cell = page.cell(row_index, column)
                    if isinstance(value, str):
                        _text(cell, value)
                    else:
                        cell.value = value
                        cell.number_format = "0.00"
                    cell.alignment = record_alignment
                    cell.border = border
                    cell.font = record_font
                page.row_dimensions[row_index].height = max(26, min(160, 14 * max(2, math.ceil(len(str(values[6])) / 55))))
                if entry is not None:
                    metadata.append((entry.id, entry.revision, entry.day.isoformat(),
                        entry.started_at.isoformat() if entry.started_at else "", entry.ended_at.isoformat() if entry.ended_at else "",
                        entry.break_hours, entry.work_type.value, entry.context.project_id or "", entry.context.work_item_id or ""))
                row_index += 1
            page.auto_filter.ref = f"A6:K{row_index - 1}"
            page.freeze_panes = "A7"
            for caption, value in _summary(sheet):
                row_index += 1
                page.cell(row_index, 7, caption).font = Font(name="Arial", bold=True)
                page.cell(row_index, 8, value).number_format = "0.00"
            row_index += 2
            page.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=11)
            page.cell(row_index, 1, _("Recorded totals at export. Editing this file does not change WorkLogger records."))
            page.cell(row_index, 1).alignment = Alignment(wrap_text=True)
            widths = (13, 9, 9, 17, 22, 25, 60, 14, 14, 14, 15)
            for column, width in enumerate(widths, 1):
                page.column_dimensions[get_column_letter(column)].width = width
            page.sheet_properties.pageSetUpPr.fitToPage = True
            page.page_setup.orientation = "landscape"
            page.page_setup.paperSize = page.PAPERSIZE_A4
            page.page_setup.fitToWidth, page.page_setup.fitToHeight = 1, 0
            page.print_title_rows = "1:6"
            page.print_area = f"A1:K{row_index}"
            sources = workbook.create_sheet(_("Record references")[:31])
            sources.append([_("Record ID"), _("Revision"), _("Date"), _("Started at"), _("Ended at"),
                            _("Historical break deduction"), _("Work type"), _("Project ID"), _("Work item ID")])
            for values in metadata:
                sources.append(values)
            sources.freeze_panes = "A2"
            for column in range(1, 10):
                sources.column_dimensions[get_column_letter(column)].width = 25
            with atomic_destination(Path(destination)) as temporary:
                workbook.save(temporary)
            workbook.close()
        except ValueError:
            return Result.failure(ValidationError("timesheet_content_invalid", "timesheet_content_invalid"))
        except Exception:
            return Result.failure(InfrastructureError("timesheet_export_failed", "timesheet_export_failed"))
        return Result.success(Path(destination))


def _text(cell, value):
    if len(value) > 32767:
        raise ValueError("timesheet_content_invalid")
    cell.value = value
    cell.data_type = "s"


class TimesheetPdfExporter:
    def export_timesheet(self, destination, sheet):
        try:
            from PySide6.QtCore import QMarginsF
            from PySide6.QtGui import QFont, QGuiApplication, QPageLayout, QPageSize, QPdfWriter, QTextDocument
            if QGuiApplication.instance() is None:
                raise RuntimeError("pdf_application_required")
            html = ["<html><body>", "<h1>" + escape(_("Timesheet")) + "</h1>",
                "<p>" + escape(sheet.start.isoformat() + " - " + sheet.end.isoformat()) + "<br>"
                + escape(_("Prepared by: {name}").format(name=sheet.author)) + "<br>"
                + escape(_("Generated at: {time}").format(time=sheet.generated_at.isoformat(timespec="seconds"))) + "</p>"]
            headings = (_("Date"), _("Time"), _("Work type"), _("Project / work item"),
                        _("Worked hours"), _("Rest hours"), _("Leave hours"), _("Content"))
            widths = (10, 12, 10, 16, 7, 7, 7, 31)
            html.append('<table border="1" cellspacing="0" cellpadding="3" width="100%"><thead><tr bgcolor="#e8f0fe">')
            html.extend(f'<th width="{width}%">' + escape(value) + "</th>" for width, value in zip(widths, headings))
            html.append("</tr></thead><tbody>")
            for values, entry in _rows(sheet):
                status = values[10]
                content = values[6] + ("\n" + status if status == _("Scheduled") else "")
                cells = (values[0], values[1] + " - " + values[2] if entry and entry.has_times else "", values[3],
                         entry.context.label if entry else "", f"{values[7]:.2f}", f"{values[8]:.2f}", f"{values[9]:.2f}", content)
                html.append("<tr>" + "".join("<td>" + escape(str(value)).replace("\n", "<br>") + "</td>" for value in cells) + "</tr>")
            html.append("</tbody></table><p>" + " | ".join(escape(caption) + f": {value:.2f}h" for caption, value in _summary(sheet)) + "</p>")
            if any(sheet.record_status(entry) == "historical" for entry in sheet.entries):
                html.append("<p>" + escape(_("Historical clock records do not contain timezone offsets.")) + "</p>")
            html.append("<p>" + escape(_("Recorded totals at export. Editing this file does not change WorkLogger records.")) + "</p></body></html>")
            with atomic_destination(Path(destination)) as temporary:
                writer = QPdfWriter(str(temporary))
                writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
                writer.setPageOrientation(QPageLayout.Orientation.Landscape)
                writer.setPageMargins(QMarginsF(12, 12, 12, 12))
                writer.setResolution(96)
                writer.setTitle(_("Timesheet"))
                writer.setCreator("WorkLogger")
                document = QTextDocument()
                font = QFont(QGuiApplication.font())
                font.setPointSizeF(9)
                document.setDefaultFont(font)
                document.setDefaultStyleSheet("body { color: #1e2035; background-color: #ffffff; } table { font-size: 8pt; } th { text-align: left; }")
                document.setHtml("".join(html))
                document.print_(writer)
                del writer
                with temporary.open("rb") as handle:
                    if handle.read(5) != b"%PDF-":
                        raise RuntimeError("timesheet_export_failed")
                    handle.seek(max(0, temporary.stat().st_size - 4096))
                    if b"%%EOF" not in handle.read(4096):
                        raise RuntimeError("timesheet_export_failed")
        except Exception:
            return Result.failure(InfrastructureError("timesheet_export_failed", "timesheet_export_failed"))
        return Result.success(Path(destination))
