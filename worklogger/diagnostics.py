"""Isolated installed-application checks for recording, recovery and delivery."""

from dataclasses import replace
from datetime import date
from pathlib import Path
import tempfile


def require_value(result):
    if not result.ok:
        raise RuntimeError(result.error.code if result.error else "operation_failed")
    return result.value


def check_workflows():
    from worklogger.bootstrap import DesktopRuntimeConfig, build_desktop_runtime
    from worklogger.domain.projects.models import WorkContext
    from worklogger.infrastructure.backup import SQLiteBackupService
    from worklogger.infrastructure.database import MigrationRunner, SQLiteConnectionFactory
    from worklogger.infrastructure.database.migrations.runner import MIGRATION_MODULES
    from worklogger.infrastructure.database.upgrade import upgrade_database
    from worklogger.infrastructure.repositories.auth_sqlite import SQLiteAuthRepository
    from worklogger.infrastructure.repositories.worklog_sqlite import SQLiteWorkLogRepository
    from worklogger.infrastructure.repositories.report_sqlite import SQLiteReportRepository
    from worklogger.infrastructure.security import PBKDF2PasswordHasher
    from worklogger.app.use_cases.project_analytics import ProjectAnalyticsHandler
    from worklogger.infrastructure.repositories.settings_sqlite import SQLiteSettingsRepository
    from openpyxl import load_workbook
    from PySide6.QtPdf import QPdfDocument

    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        factory = SQLiteConnectionFactory(root / "previous.db")
        MigrationRunner(factory, MIGRATION_MODULES[:1]).run_pending()
        auth = SQLiteAuthRepository(factory, password_hasher=PBKDF2PasswordHasher(iterations=1000))
        previous = auth.create_user("sample", "synthetic-password", recovery_key=None, is_admin=True)
        day = date(2026, 5, 21)
        with factory.transaction() as connection:
            connection.execute('INSERT INTO worklog(user_id,d,start,end,"break",note) VALUES(?,?,?,?,?,?)',
                               (previous.id, day.isoformat(), "09:00", "12:00", 0.5, "Historical work"))
            connection.execute("INSERT INTO reports(user_id,type,period_start,period_end,content,created_at) VALUES(?,'daily',?,?,?,datetime('now'))",
                               (previous.id, day.isoformat(), day.isoformat(), "Historical report"))
        upgraded = root / "worklog.db"
        summary = upgrade_database(Path(factory.database_path), upgraded)
        if summary.time_records != 1 or summary.reports != 1:
            raise RuntimeError("upgrade_verification_failed")
        runtime = require_value(build_desktop_runtime(DesktopRuntimeConfig(database_path=upgraded,
            user_id=previous.id, minimal_mode=False, password_iterations=1000), argv=[]))
        try:
            current_auth = SQLiteAuthRepository(runtime.connection_factory, password_hasher=PBKDF2PasswordHasher(iterations=1000))
            if current_auth.verify_user("sample", "synthetic-password") is None:
                raise RuntimeError("credential_verification_failed")
            service = runtime.window.entry_panel.view_model.service
            project = require_value(service.projects.save_project("Research"))
            item = require_value(service.projects.save_work_item(project.id, "Review"))
            context = WorkContext(project.id, item.id)
            entry = require_value(service.save_manual(day, "13:00", "15:00", "normal", "Work content", context=context))
            history = service.repository.list_for_day(previous.id, day)[0]
            assigned = require_value(service.associate((history,), context))[0]
            require_value(service.undo(service.repository.latest_change(previous.id).id))
            if assigned.break_hours != 0.5 or entry.worked_hours() != 2:
                raise RuntimeError("accounting_verification_failed")
            analytics = ProjectAnalyticsHandler(service.repository, SQLiteSettingsRepository(runtime.connection_factory))
            if sum(row.work_hours for row in require_value(analytics.handle(previous.id, day, day))) != 4.5:
                raise RuntimeError("analytics_verification_failed")
            reports = runtime.window.reports_page._view_model
            state = require_value(reports.load("daily", day))
            generated = require_value(reports.generate_with_sources(state))
            state = replace(state, provenance=generated.provenance)
            saved = require_value(reports.save(state, generated.content))
            updated = require_value(reports.save(saved, "Updated report"))
            restored = require_value(reports.restore_revision(updated, saved.revision))
            if restored.content != generated.content or len(require_value(reports.list_revisions(restored))) != 4:
                raise RuntimeError("report_recovery_verification_failed")
            require_value(reports.export_timesheet(root / "timesheet.xlsx", day))
            workbook = load_workbook(root / "timesheet.xlsx", read_only=True)
            workbook.close()
            require_value(reports.export_timesheet(root / "timesheet.pdf", day, format="pdf"))
            document = QPdfDocument()
            if document.load(str(root / "timesheet.pdf")) != QPdfDocument.Error.None_ or document.pageCount() < 1:
                raise RuntimeError("pdf_verification_failed")
            document.close()
            del document
            backup = SQLiteBackupService(runtime.connection_factory, expected_username="sample", requesting_user_id=previous.id)
            require_value(backup.backup_database(root / "backup.db"))
            backup_factory = SQLiteConnectionFactory(root / "backup.db")
            if len(SQLiteWorkLogRepository(backup_factory).list_for_day(previous.id, day)) != 2:
                raise RuntimeError("backup_verification_failed")
            if SQLiteReportRepository(backup_factory).get_for_period(previous.id, "daily", day, day).content != restored.content:
                raise RuntimeError("backup_verification_failed")
        finally:
            if runtime.job_runner is not None:
                runtime.job_runner.shutdown(wait=True)
            if runtime.local_inference is not None:
                runtime.local_inference.close()
            runtime.window.close()
            runtime.window.deleteLater()
