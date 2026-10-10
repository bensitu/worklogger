from dataclasses import replace
from datetime import date, datetime, timezone
import os
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication, QMessageBox

from tests.app.test_notes_and_reports_use_cases import (
    MemoryCalendarRepository, MemoryQuickLogRepository, MemoryTemplateProvider,
    MemoryTemplateRepository, MemoryWorkLogRepository,
)
from worklogger.app.commands.report_commands import DeleteReportCommand, SaveReportCommand
from worklogger.app.use_cases.ai import RewriteTextHandler
from worklogger.app.use_cases.reports import (
    DeleteReportHandler, GenerateReportHandler, GetReportForPeriodHandler, ListReportsHandler,
    ResetReportTemplateHandler, SaveReportHandler, SaveReportTemplateHandler, ReportRevisionService,
)
from worklogger.domain.reporting.models import Report
from worklogger.domain.reporting.periods import daily_period, monthly_period, weekly_period
from worklogger.infrastructure.database import MigrationRunner, SQLiteConnectionFactory
from worklogger.infrastructure.export import MarkdownExporter
from worklogger.infrastructure.i18n import _, set_language
from worklogger.infrastructure.repositories import SQLiteAuthRepository, SQLiteReportRepository
from worklogger.infrastructure.security import PBKDF2PasswordHasher
from worklogger.presentation.reporting.dialog import ReportDialog
from worklogger.presentation.shell.pages import ReportsPage
from worklogger.presentation.theme import install_bundled_fonts
from worklogger.presentation.viewmodels.reports import ReportEditorViewModel
from worklogger.presentation.job_runner import ImmediateJobRunner, QtJobRunner


class ReportHistoryTests(unittest.TestCase):
    def test_timesheet_exports_preserve_periods_unicode_and_report_drafts(self):
        from openpyxl import load_workbook
        from PySide6.QtPdf import QPdfDocument
        import unicodedata
        from worklogger.domain.worklog.models import WorkLog, WorkType
        from worklogger.domain.projects.models import WorkContext
        from worklogger.infrastructure.repositories.worklog_sqlite import SQLiteWorkLogRepository
        from worklogger.infrastructure.export.timesheets import TimesheetXlsxExporter, TimesheetPdfExporter
        from worklogger.app.use_cases.timesheets import ExportTimesheetHandler
        records = SQLiteWorkLogRepository(self.factory)
        original = records.save_entry(WorkLog(self.user.id, self.day, "09:00", "12:00", 0.5, "=SUM(A1:A2)\n\u5de5\u4f5c\u5185\u5bb9",
            context=WorkContext(project_label="Research", work_item_label="Review")))
        records.save_entry(WorkLog(self.user.id, self.day, "12:00", "13:00", work_type=WorkType.BREAK))
        records.save_entry(WorkLog(self.user.id, self.day, "14:00", "16:00", note="Afternoon"))
        records.save_entry(WorkLog(self.other_user.id, self.day, "09:00", "10:00", note="Other account"))
        handler = ExportTimesheetHandler(records=records, exporters={"xlsx": TimesheetXlsxExporter(), "pdf": TimesheetPdfExporter()}, clock=lambda: self.stamp)
        self.model._timesheet_export_handler = handler
        page = self.page()
        page.editor.setPlainText("Unsaved report draft")
        target = Path(self.factory.database_path).parent / "timesheet.xlsx"
        with patch("worklogger.presentation.shell.reports_page.QFileDialog.getSaveFileName", return_value=(str(target), "")):
            page._choose_export_scope("timesheet_day_xlsx")
        self.assertTrue(target.is_file())
        book = load_workbook(target, data_only=False)
        try:
            sheet = book.worksheets[0]
            self.assertEqual(sheet["G7"].value, original.note)
            self.assertEqual(sheet["G7"].data_type, "s")
            self.assertEqual([sheet.cell(row, 2).value for row in range(7, 10)], ["09:00", "12:00", "14:00"])
            self.assertEqual(sum(sheet.cell(row, 8).value for row in range(7, 10)), 4.5)
            self.assertEqual(sum(sheet.cell(row, 9).value for row in range(7, 10)), 1.5)
            self.assertEqual(book.worksheets[1].max_row, 4)
            self.assertNotIn("Other account", str(tuple(sheet.values)))
        finally:
            book.close()
        pdf = target.with_suffix(".pdf")
        self.assertTrue(self.model.export_timesheet(pdf, self.day, whole_month=True, format="pdf").ok)
        document = QPdfDocument()
        self.assertEqual(document.load(str(pdf)), QPdfDocument.Error.None_)
        try:
            text = "\n".join(document.getAllText(index).text() for index in range(document.pageCount()))
            self.assertIn("Afternoon", text)
            self.assertIn("\u5de5\u4f5c\u5185\u5bb9", "".join(unicodedata.normalize("NFKC", text).split()))
            self.assertNotIn("Other account", text)
            self.assertGreater(document.pageCount(), 0)
        finally:
            document.close()
        self.assertEqual(page.editor.toPlainText(), "Unsaved report draft")
        self.assertEqual(self.repository.list_by_type(self.user.id, "daily"), ())
        self.assertEqual(records.get_entry(self.user.id, original.id), original)

    def test_timesheet_export_failures_preserve_destinations_and_classify_future_rest(self):
        from worklogger.domain.reporting.timesheets import Timesheet
        from worklogger.domain.worklog.models import WorkLog, WorkType
        from worklogger.infrastructure.export.timesheets import TimesheetXlsxExporter
        from datetime import timedelta
        start = self.stamp.replace(hour=12, minute=0, second=0)
        end = start + timedelta(hours=1)
        rest = WorkLog(self.user.id, start.date(), "12:00", "13:00", work_type=WorkType.BREAK,
                       started_at=start, ended_at=end)
        snapshot = Timesheet(self.day, self.day, "Mary", start + timedelta(minutes=5), 8, (rest,))
        self.assertEqual(snapshot.record_status(rest), "scheduled")
        self.assertEqual(snapshot.statistics.total_hours, 0)
        self.assertEqual(snapshot.rest_hours, 1)
        target = Path(self.factory.database_path).parent / "preserved.xlsx"
        target.write_bytes(b"Existing file")
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("synthetic failure")):
            self.assertFalse(TimesheetXlsxExporter().export_timesheet(target, snapshot).ok)
        self.assertEqual(target.read_bytes(), b"Existing file")
        oversized = replace(rest, note="X" * 32768)
        invalid = replace(snapshot, entries=(oversized,))
        self.assertFalse(TimesheetXlsxExporter().export_timesheet(target, invalid).ok)
        self.assertEqual(target.read_bytes(), b"Existing file")

    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])
        install_bundled_fonts()

    def setUp(self):
        directory = self.enterContext(tempfile.TemporaryDirectory())
        self.factory = SQLiteConnectionFactory(Path(directory) / "reports.db")
        MigrationRunner(self.factory).run_pending()
        auth = SQLiteAuthRepository(self.factory, password_hasher=PBKDF2PasswordHasher(iterations=1000))
        self.user = auth.create_user("report.user", "test-password", recovery_key=None, is_admin=False)
        self.other_user = auth.create_user("other.user", "test-password", recovery_key=None, is_admin=False)
        self.repository = SQLiteReportRepository(self.factory)
        templates = MemoryTemplateRepository()
        self.model = ReportEditorViewModel(
            user_id=self.user.id,
            generate_handler=GenerateReportHandler(
                work_logs=MemoryWorkLogRepository(()), quick_logs=MemoryQuickLogRepository(()),
                calendar_events=MemoryCalendarRepository(()), templates=MemoryTemplateProvider(),
            ),
            get_report_handler=GetReportForPeriodHandler(self.repository),
            save_report_handler=SaveReportHandler(self.repository),
            list_reports_handler=ListReportsHandler(self.repository),
            delete_report_handler=DeleteReportHandler(self.repository),
            revision_service=ReportRevisionService(self.repository),
            save_template_handler=SaveReportTemplateHandler(templates),
            reset_template_handler=ResetReportTemplateHandler(templates),
            markdown_exporter=MarkdownExporter(), rewrite_handler=RewriteTextHandler(),
        )
        self.information = self.enterContext(patch.object(QMessageBox, "information"))
        self.warning = self.enterContext(patch.object(QMessageBox, "warning"))
        self.day = date(2026, 5, 21)
        self.stamp = datetime(2026, 5, 21, 6, 7, 8, tzinfo=timezone.utc)
        set_language("en_US")

    def tearDown(self):
        set_language("en_US")

    def seed(self, report_type, content, *, user_id=None):
        period = {"daily": daily_period(self.day), "weekly": weekly_period(self.day),
                  "monthly": monthly_period(self.day.year, self.day.month)}[report_type]
        return self.repository.save(Report(None, user_id or self.user.id, report_type,
                                           period.start, period.end, content, self.stamp))

    def test_report_recovery_preserves_sources_versions_and_account_boundaries(self):
        saved = self.seed("daily", "Original")
        state = self.model.load("daily", self.day).value
        changed = self.model.save(state, "Updated").value
        self.assertEqual(changed.revision, 1)
        conflict = self.model.save(state, "Stale edit")
        self.assertFalse(conflict.ok)
        self.assertEqual(conflict.error.code, "report_conflict")
        self.assertEqual([value.content for value in self.model.list_revisions(changed).value], ["Updated", "Original"])
        restored = self.model.restore_revision(changed, 0).value
        self.assertEqual((restored.report_id, restored.content, restored.revision), (saved.id, "Original", 2))
        self.assertEqual([value.revision for value in self.model.list_revisions(restored).value], [2, 1, 0])
        self.assertEqual(self.repository.list_revisions(self.other_user.id, saved.id), ())
        with self.assertRaisesRegex(ValueError, "report_not_found"):
            self.repository.restore_revision(self.other_user.id, saved.id, 0, 2)
        page = self.page()
        page.editor.setPlainText("Unsaved")
        from worklogger.presentation.widgets.report_revisions import ReportRevisionsDialog
        dialog = ReportRevisionsDialog(self.model, restored, job_runner=ImmediateJobRunner())
        dialog._load()
        self.assertEqual(dialog.versions.count(), 3)
        dialog.versions.setCurrentRow(1)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            dialog._restore()
        self.assertEqual(page.editor.toPlainText(), "Unsaved")
        self.assertEqual(self.model.load("daily", self.day).value.revision, 2)
        dialog.close()
        self.repository.remove(self.user.id, saved.id)
        with self.factory.connection() as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM report_revisions").fetchone()[0], 0)

    def test_generation_sources_and_retained_versions_are_bounded(self):
        draft = self.model.load("daily", self.day).value
        self.assertTrue(draft.provenance.generated_at)
        self.assertEqual(draft.provenance.language, "en_US")
        saved = self.model.save(draft, "First").value
        self.assertEqual(saved.provenance, draft.provenance)
        for index in range(55):
            saved = self.model.save(saved, f"Version {index}").value
        versions = self.model.list_revisions(saved).value
        self.assertEqual(len(versions), 50)
        self.assertEqual(versions[0].revision, 55)
        self.assertEqual(versions[-1].revision, 6)
        before = saved
        self.assertFalse(self.model.save(saved, "X" * (1024 * 1024 + 1)).ok)
        self.assertEqual(self.model.load("daily", self.day).value, before)

    def test_existing_reports_migrate_with_recoverable_content(self):
        from worklogger.infrastructure.database.migrations.runner import MIGRATION_MODULES
        path = Path(self.factory.database_path).parent / "existing.db"
        factory = SQLiteConnectionFactory(path)
        MigrationRunner(factory, MIGRATION_MODULES[:12]).run_pending()
        auth = SQLiteAuthRepository(factory, password_hasher=PBKDF2PasswordHasher(iterations=1000))
        user = auth.create_user("migration.user", "synthetic-password", recovery_key=None, is_admin=False)
        with factory.transaction(write=True) as connection:
            connection.execute("INSERT INTO reports(user_id,type,period_start,period_end,content,created_at) VALUES(?,'daily',?,?,?,?)",
                (user.id, self.day.isoformat(), self.day.isoformat(), "Existing content", self.stamp.isoformat()))
        self.assertEqual(MigrationRunner(factory).run_pending(), (13,))
        self.assertEqual(MigrationRunner(factory).run_pending(), ())
        repository = SQLiteReportRepository(factory)
        report = repository.get_for_period(user.id, "daily", self.day, self.day)
        self.assertEqual(repository.list_revisions(user.id, report.id)[0].content, "Existing content")
        self.assertEqual(report.provenance.generated_at, "")
        self.assertTrue(list(path.parent.glob("*upgrade*")))

    def page(self, report_type="daily"):
        page = ReportsPage(self.model, self.day, job_runner=ImmediateJobRunner())
        self.assertTrue(page.refresh())
        page.report_type_control.set_value(report_type)
        self.addCleanup(page.deleteLater)
        self.addCleanup(page.close)
        return page

    def test_history_delete_confirms_and_updates_only_the_target_report(self):
        for report_type in ("daily", "weekly", "monthly"):
            first = self.seed(report_type, "First")
            current = self.seed(report_type, "Current")
            page = self.page(report_type)
            page.editor.setPlainText("Unsubmitted edits")
            button = next(button for button in page.history_panel._buttons.values() if button.property("report_id") == first.id)
            with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
                button.delete_button.click()
            self.assertEqual(len(self.repository.list_by_type(self.user.id, report_type)), 2)
            with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
                button.delete_button.click()
            self.assertEqual([report.id for report in self.repository.list_by_type(self.user.id, report_type)], [current.id])
            self.assertEqual(page.editor.toPlainText(), "Unsubmitted edits")
            self.assertEqual(page._states[report_type].report_id, current.id)
            active = next(iter(page.history_panel._buttons.values()))
            with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
                active.delete_button.click()
            self.assertEqual(self.repository.list_by_type(self.user.id, report_type), ())
            self.assertIsNone(page._states[report_type].report_id)
            self.assertEqual(page.editor.toPlainText(), "")
            self.assertFalse(page.has_unsaved_changes)
            page.editor.setPlainText("Replacement")
            page._save_current()
            saved = self.repository.list_by_type(self.user.id, report_type)
            self.assertEqual(len(saved), 1)
            self.assertNotEqual(saved[0].id, current.id)
            page._confirm_discard = lambda: True

    def test_history_delete_rejects_stale_foreign_and_failed_storage_operations(self):
        saved = self.seed("daily", "Original")
        page = self.page()
        stale = page.history_panel._items[0]
        page.editor.setPlainText("Unsubmitted edit")
        self.repository.save(replace(saved, content="Changed elsewhere"))
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            page._delete_history_item(stale)
        self.assertEqual(page.last_error.code, "report_conflict")
        self.assertEqual(page.editor.toPlainText(), "Unsubmitted edit")
        self.assertEqual(len(self.repository.list_by_type(self.user.id, "daily")), 1)
        handler = DeleteReportHandler(self.repository)
        rejected = handler.handle(DeleteReportCommand(self.other_user.id, saved.id, "Changed elsewhere"))
        self.assertEqual(rejected.error.code, "report_not_found")
        item = replace(stale, content="Changed elsewhere")
        for failure in (OSError("storage unavailable"), ValueError("invalid stored data")):
            with patch.object(self.repository, "remove", side_effect=failure), patch.object(
                QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
                page._delete_history_item(item)
            self.assertEqual(page.last_error.code, "report_delete_failed")
            self.assertEqual(page.editor.toPlainText(), "Unsubmitted edit")
            self.assertFalse(page._delete_busy)
            self.assertTrue(page.history_panel.isEnabled())
        self.assertFalse(self.model.delete(replace(item, user_id=self.other_user.id)).ok)
        page._confirm_discard = lambda: True

    def test_background_delete_serializes_editor_actions_without_blocking_events(self):
        self.seed("daily", "Original")
        page = self.page()
        page._job_runner = QtJobRunner(page)
        self.addCleanup(lambda: page._job_runner.shutdown(wait=True))
        item = page.history_panel._items[0]
        release = threading.Event()
        threads = []
        remove = self.repository.remove

        def delayed(*args, **kwargs):
            threads.append(threading.get_ident())
            release.wait(5)
            return remove(*args, **kwargs)

        with patch.object(self.repository, "remove", side_effect=delayed), patch.object(
            QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            try:
                page._delete_history_item(item)
                self.assertTrue(page._delete_busy)
                self.assertFalse(page.confirm_leave())
                self.assertTrue(page.editor.isReadOnly())
                self.assertFalse(page.history_panel.isEnabled())
                self.assertFalse(page.save_button.isEnabled())
                self.app.processEvents()
                self.assertEqual(page.editor.toPlainText(), "Original")
            finally:
                release.set()
            deadline = time.monotonic() + 5
            while page.is_busy and time.monotonic() < deadline:
                self.app.processEvents()
                time.sleep(0.01)
        self.assertFalse(page._delete_busy)
        self.assertNotEqual(threads[0], threading.get_ident())
        self.assertTrue(page.history_panel.isEnabled())
        self.assertFalse(page.editor.isReadOnly())
        self.assertEqual(page.editor.toPlainText(), "")
        self.assertEqual(self.repository.list_by_type(self.user.id, "daily"), ())

    def test_confirmed_overwrite_changes_only_the_selected_report(self):
        for report_type in ("daily", "weekly", "monthly"):
            original = self.seed(report_type, "Original")
            newer = self.seed(report_type, "Another report")
            page = self.page(report_type)
            self.assertEqual(page._states[report_type].report_id, newer.id)
            item = next(item for item in page.history_panel._items if item.report_id == original.id)
            page._select_history_item(item)
            self.assertEqual(page._states[report_type].report_id, original.id)
            self.assertEqual(sum(bool(button.property("active")) for button in page.history_panel._buttons.values()), 1)
            page.editor.setPlainText("Edited")
            with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No) as question:
                page.save_button.click()
                self.assertEqual(question.call_args.args[-1], QMessageBox.StandardButton.No)
            self.assertEqual(page.editor.toPlainText(), "Edited")
            self.assertTrue(page.has_unsaved_changes)
            self.assertEqual(page._states[report_type].content, "Original")
            self.assertEqual(len(self.repository.list_by_type(self.user.id, report_type)), 2)
            with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
                page.save_button.click()
            reports = {report.id: report for report in self.repository.list_by_type(self.user.id, report_type)}
            self.assertEqual(len(reports), 2)
            self.assertEqual(reports[original.id].content, "Edited")
            self.assertEqual(reports[original.id].created_at, original.created_at)
            self.assertEqual(reports[newer.id], newer)
            self.assertEqual(page._states[report_type].report_id, original.id)
            self.assertFalse(page.has_unsaved_changes)
            self.assertEqual([button.property("report_id") for button in page.history_panel._buttons.values()
                              if button.property("active")], [original.id])
            with patch.object(QMessageBox, "question") as question:
                page.save_button.click()
                question.assert_not_called()
            self.assertEqual(len(self.repository.list_by_type(self.user.id, report_type)), 2)

    def test_new_save_retains_identity_for_subsequent_edits(self):
        page = self.page()
        self.assertIsNone(page._states["daily"].report_id)
        page.editor.setPlainText("New report")
        with patch.object(QMessageBox, "question") as question:
            page.save_button.click()
            question.assert_not_called()
        state = page._states["daily"]
        self.assertIsNotNone(state.report_id)
        self.assertIsNotNone(state.created_at)
        page.editor.setPlainText("Updated report")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            page.save_button.click()
        saved = self.repository.list_by_type(self.user.id, "daily")
        self.assertEqual(len(saved), 1)
        self.assertEqual(saved[0].id, state.report_id)
        self.assertEqual(saved[0].created_at, state.created_at)
        self.assertEqual(saved[0].content, "Updated report")

    def test_missing_report_is_not_recreated_and_keeps_the_draft(self):
        saved = self.seed("daily", "Original")
        page = self.page()
        self.repository.remove(self.user.id, saved.id)
        page.editor.setPlainText("Edited")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            page.save_button.click()
        self.assertEqual(page.last_error.code, "report_not_found")
        self.assertEqual(self.repository.list_by_type(self.user.id, "daily"), ())
        self.assertEqual(page.editor.toPlainText(), "Edited")
        self.assertTrue(page.has_unsaved_changes)
        self.warning.assert_called_once()
        self.assertIn("no longer available", self.warning.call_args.args[2])

    def test_storage_failure_keeps_saved_content_and_editor_changes(self):
        saved = self.seed("daily", "Original")
        page = self.page()
        page.editor.setPlainText("Edited")
        with patch.object(self.repository, "save", side_effect=RuntimeError("Storage unavailable")):
            with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
                page.save_button.click()
        self.assertEqual(page.last_error.code, "report_save_failed")
        self.assertEqual(self.repository.list_by_type(self.user.id, "daily"), (saved,))
        self.assertEqual(page.editor.toPlainText(), "Edited")
        self.assertTrue(page.has_unsaved_changes)
        self.assertIn("changes have been retained", self.warning.call_args.args[2])

    def test_update_rejects_other_accounts_types_and_periods(self):
        saved = self.seed("daily", "Original")
        command = SaveReportCommand(self.user.id, "daily", saved.period_start, saved.period_end,
                                    "Edited", report_id=saved.id)
        invalid = (
            replace(command, user_id=self.other_user.id),
            replace(command, report_type="weekly", period_start=date(2026, 5, 18), period_end=date(2026, 5, 24)),
            replace(command, period_start=date(2026, 5, 22), period_end=date(2026, 5, 22)),
            replace(command, report_id=saved.id + 100),
        )
        for value in invalid:
            result = SaveReportHandler(self.repository).handle(value)
            self.assertFalse(result.ok)
            self.assertEqual(result.error.code, "report_not_found")
        foreign = self.seed("daily", "Other account", user_id=self.other_user.id)
        state = self.model.load("daily", self.day).value
        result = self.model.save(replace(state, user_id=self.other_user.id, report_id=foreign.id), "Edited")
        self.assertFalse(result.ok)
        self.assertEqual(self.repository.list_by_type(self.user.id, "daily"), (saved,))
        self.assertEqual(self.repository.list_by_type(self.other_user.id, "daily"), (foreign,))

    def test_report_dialog_confirms_existing_report_updates(self):
        saved = self.seed("weekly", "Original")
        dialog = ReportDialog(self.model, self.day)
        self.addCleanup(dialog.deleteLater)
        self.assertTrue(dialog.refresh())
        dialog.tabs.setCurrentIndex(1)
        dialog.weekly_editor.setPlainText("Edited")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            dialog.save_button.click()
        self.assertEqual(self.repository.list_by_type(self.user.id, "weekly"), (saved,))
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            dialog.save_button.click()
        reports = self.repository.list_by_type(self.user.id, "weekly")
        self.assertEqual(len(reports), 1)
        self.assertEqual(reports[0].id, saved.id)
        self.assertEqual(reports[0].content, "Edited")

    def test_history_timestamps_identifiers_selection_and_search_are_localized(self):
        for report_type in ("daily", "weekly", "monthly"):
            self.seed(report_type, "Original")
            self.seed(report_type, "Another report")
        for language in ("en_US", "ja_JP", "ko_KR", "zh_CN", "zh_TW"):
            set_language(language)
            created_label = _("First saved: {time}").split("{time}")[0]
            report_label = _("Report #{report_id}").split("{report_id}")[0]
            page = self.page()
            for report_type in ("daily", "weekly", "monthly"):
                page.report_type_control.set_value(report_type)
                buttons = tuple(page.history_panel._buttons.values())
                self.assertEqual(len(buttons), 2)
                self.assertNotEqual(buttons[0].text(), buttons[1].text())
                for button in buttons:
                    self.assertIn(created_label, button.text())
                    self.assertIn(self.stamp.astimezone().strftime("%Y-%m-%d %H:%M:%S"), button.text())
                    self.assertIn(report_label, button.text())
                buttons[1].click()
                selected = page._states[report_type].report_id
                self.assertEqual(sum(bool(button.property("active")) for button in buttons), 1)
                page.history_panel.search_line_edit.setText(f"#{selected}")
                self.assertEqual(len(page.history_panel._buttons), 1)
                self.assertTrue(page.history_panel._buttons[0].property("active"))
                page.history_panel.search_line_edit.clear()
                self.assertEqual(sum(bool(button.property("active")) for button in page.history_panel._buttons.values()), 1)

    def test_saved_daily_exports_use_latest_owned_report_per_day_without_changing_editor(self):
        older = self.seed("daily", "Older daily")
        latest = self.seed("daily", "Latest daily")
        self.seed("monthly", "Monthly report is not a daily report")
        self.seed("daily", "Another account", user_id=self.other_user.id)
        previous_day = self.day.replace(day=20)
        self.repository.save(Report(None, self.user.id, "daily", previous_day, previous_day,
                                    "Previous day", self.stamp))
        next_month = self.day.replace(month=6)
        self.repository.save(Report(None, self.user.id, "daily", next_month, next_month, "Next month", self.stamp))
        destination = Path(self.factory.database_path).parent / "daily.md"
        day_result = self.model.export_saved_daily(destination, self.day)
        self.assertTrue(day_result.ok, day_result.error)
        exported = destination.read_text(encoding="utf-8")
        self.assertIn("Latest daily", exported)
        self.assertNotIn("Older daily", exported)
        self.assertNotIn("Another account", exported)
        self.assertIn(f"#{latest.id}", exported)
        self.assertNotIn(f"#{older.id}", exported)
        month_result = self.model.export_saved_daily(destination, self.day, whole_month=True)
        self.assertTrue(month_result.ok, month_result.error)
        exported = destination.read_text(encoding="utf-8")
        self.assertLess(exported.index("Previous day"), exported.index("Latest daily"))
        self.assertNotIn("Monthly report", exported)
        self.assertNotIn("Next month", exported)
        self.assertFalse(self.model.export_saved_daily(destination, self.day.replace(month=7)).ok)
        self.assertEqual(destination.read_text(encoding="utf-8"), exported)


if __name__ == "__main__":
    unittest.main()
